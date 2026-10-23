"""END-008 integration package (END008-BatteryHealthMonitoring/star_integration): phase 1, no STAR backend involved.

Writer, base builder and reader are tested on a small fixture (tests/fixtures/end008). The tests that need the END-008
delivery (converter, plot script, executable) are skipped when the folder is missing (it is not in git, the CI has no copy).
"""
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
VENDOR = ROOT / "END008-BatteryHealthMonitoring"
PKG = VENDOR / "star_integration"
FIX = Path(__file__).resolve().parent / "fixtures" / "end008"
if not PKG.is_dir():
    raise unittest.SkipTest("END008-BatteryHealthMonitoring/star_integration not present")
if str(PKG) not in sys.path:
    sys.path.insert(0, str(PKG))

import star_end008_input as inp          # noqa: E402
import star_end008_output as out         # noqa: E402
import star_end008_power_budget as pb    # noqa: E402

try:
    import openpyxl                      # noqa: E402
except ImportError:                      # pragma: no cover
    openpyxl = None
HAS_VENDOR = (VENDOR / inp.CONVERTER).is_file() and (VENDOR / out.PLOTTER).is_file()
HAS_EXE = (VENDOR / inp.EXE).is_file()
SUMMARY = FIX / "power_budget_summary_small.json"


class TestAllocateOrbits(unittest.TestCase):
    def test_largest_remainder_keeps_the_rounded_total(self):
        r = pb.allocate_orbits({"idle": 7.805, "tx_uhf": 0, "tx_sband": 5, "payload": 2})
        self.assertEqual(r, {"idle": 8, "tx_uhf": 0, "tx_sband": 5, "payload": 2})
        r = pb.allocate_orbits({"a": 2.4, "b": 2.4, "c": 2.4})
        self.assertEqual(sum(r.values()), 7)

    def test_integers_unchanged_and_negatives_ignored(self):
        self.assertEqual(pb.allocate_orbits({"a": 3, "b": -1}), {"a": 3, "b": 0})


class TestBudgetValues(unittest.TestCase):
    def test_minimum_generation_and_notes(self):
        v = pb.budget_values(SUMMARY)
        self.assertAlmostEqual(v["generated_Wh"], 21.927)
        self.assertEqual(v["battery_Wh"], 84.0)
        self.assertTrue(any("minimum" in n for n in v["notes"]))
        self.assertEqual([s["name"] for s in v["subsystems"]], ["OBDH", "TX UHF", "TX S BAND", "CAMERAS", "THERMAL"])


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class TestWorkbook(unittest.TestCase):
    def test_layout_labels(self):
        with tempfile.TemporaryDirectory() as tmp:
            pb.write_power_budget_xlsx(SUMMARY, Path(tmp) / "pb.xlsx")
            ws = openpyxl.load_workbook(Path(tmp) / "pb.xlsx").active
            self.assertEqual((ws["A1"].value, ws["B1"].value, ws["D1"].value, ws["J1"].value),
                             ("Subsystem", "Power [W]", "IDLE mode", "PAYLOAD mode"))
            self.assertEqual(ws["C2"].value, "Duty cycle - Sun/Eclipse")
            self.assertEqual((ws["N3"].value, ws["N5"].value, ws["N8"].value), ("Period [min]", "Eclipse [min]", "Margin"))
            self.assertEqual(ws["O3"].value, 97.2621)
            self.assertEqual(ws["A5"].value, "TX UHF")

    @unittest.skipUnless(HAS_VENDOR, "END-008 delivery not present")
    def test_official_converter_reads_it(self):
        with tempfile.TemporaryDirectory() as tmp:
            x = Path(tmp) / "pb.xlsx"
            v = pb.write_power_budget_xlsx(SUMMARY, x)
            subs, duty, orbit, seq = inp.load_converter().read_power_budget(x)
            self.assertEqual(list(subs), ["OBDH", "TX_UHF", "TX_S_BAND", "CAMERAS", "THERMAL"])
            self.assertEqual(set(duty), {"IDLE", "TX_UHF", "TX_S_BAND", "PAYLOAD"})
            self.assertAlmostEqual(orbit["generated"], v["generated_Wh"])
            self.assertAlmostEqual(orbit["battery"], 84.0)
            self.assertEqual(seq, [{"mode": "PAYLOAD", "orbits": 2}, {"mode": "TX_S_BAND", "orbits": 5},
                                   {"mode": "IDLE", "orbits": 8}])
            self.assertEqual(duty["PAYLOAD"]["CAMERAS"], [0.01712, 0.0])

    @unittest.skipUnless(HAS_VENDOR and (VENDOR / "examples" / "power_budget_ex1.xlsx").is_file(), "END-008 delivery not present")
    def test_same_input_as_the_delivered_example(self):
        """The delivered sheet re-expressed as an END-004 summary and written by our writer reads back identically."""
        conv = inp.load_converter()
        subs, duty, orbit, seq = conv.read_power_budget(VENDOR / "examples" / "power_budget_ex1.xlsx")
        key = {"IDLE": "idle", "TX_UHF": "tx_uhf", "TX_S_BAND": "tx_sband", "PAYLOAD": "payload"}
        per_day = {key[s["mode"]]: s["orbits"] for s in seq}
        summary = {"input_config": {
            "orbit": {"period_min": orbit["period"], "eclipse_min": orbit["eclipse"]},
            "generation": {"power_generated_per_orbit_Wh_by_mode": {"idle": orbit["generated"]}},
            "battery": {"capacity_Wh": orbit["battery"]}, "margin_fraction": orbit["margin"],
            "operations": {"mode_orbits_per_day": per_day},
            "subsystems": [{"name": n.replace("_", " "), "powerW": p,
                            "duty": {key[m]: {"sun": d.get(n, [0, 0])[0], "eclipse": d.get(n, [0, 0])[1]}
                                     for m, d in duty.items()}} for n, p in subs.items()]}}
        with tempfile.TemporaryDirectory() as tmp:
            pb.write_power_budget_xlsx(summary, Path(tmp) / "ours.xlsx")
            subs2, duty2, orbit2, seq2 = conv.read_power_budget(Path(tmp) / "ours.xlsx")
        self.assertEqual(subs2, subs)
        self.assertEqual(duty2, duty)
        self.assertEqual(orbit2, orbit)
        self.assertEqual(sorted(seq2, key=lambda s: s["mode"]), sorted(seq, key=lambda s: s["mode"]))


class TestBase(unittest.TestCase):
    def test_test_and_full_runs(self):
        tpl = inp.load_template()
        base, prov = inp.build_base(tpl, run_mode="test", battery_wh=84.0, lifetime_years=5.0, temperature_C=12.5)
        self.assertEqual(base["mission"]["duration_years"], 1.0)
        self.assertEqual(base["simulation"]["time_step_s"], 120.0)
        self.assertEqual(base["pack"]["cells_in_parallel"], 4)          # 84 / (2 x 3.6 x 2.9)
        self.assertEqual(base["pack"]["temperature_C"], 12.5)
        self.assertFalse(base["output"]["save_time_series"])
        self.assertTrue(any(k == "cell" for k, _, _ in prov))
        base, _ = inp.build_base(tpl, run_mode="full", battery_wh=84.0, lifetime_years=5.0)
        self.assertEqual((base["mission"]["duration_years"], base["simulation"]["time_step_s"]), (5.0, 60.0))
        self.assertEqual(base["pack"]["temperature_C"], inp.DEFAULT_TEMPERATURE_C)
        base, _ = inp.build_base(tpl, run_mode="full", battery_wh=84.0)
        self.assertEqual(base["mission"]["duration_years"], inp.DEFAULT_DURATION_YEARS)

    def test_template_file_overrides_defaults(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "t.json"
            p.write_text(json.dumps({"cell": {"nominal_capacity_Ah": 3.4}, "_note": "x"}), "utf-8")
            tpl = inp.load_template(p)
        self.assertEqual(tpl["cell"]["nominal_capacity_Ah"], 3.4)
        self.assertEqual(tpl["cell"]["nominal_voltage_V"], 3.6)
        self.assertNotIn("_note", tpl)

    def test_battery_temperature(self):
        th = json.loads((FIX / "thermal_summary_small.json").read_text("utf-8"))
        t, why = inp.battery_temperature_C([th, th])
        self.assertEqual(t, th["nodes"]["BAT"]["mean_C"])
        self.assertEqual(inp.battery_temperature_C([])[0], inp.DEFAULT_TEMPERATURE_C)

    def test_mission_lifetime_from_text(self):
        self.assertEqual(inp.mission_lifetime_years(["The mission lifetime shall be at least 5 years."]), 5.0)
        self.assertEqual(inp.mission_lifetime_years(["Operational lifetime of three years"]), 3.0)
        self.assertIsNone(inp.mission_lifetime_years(["no duration here"]))

    @unittest.skipUnless(HAS_VENDOR and openpyxl is not None, "END-008 delivery not present")
    def test_input_json_from_converter(self):
        with tempfile.TemporaryDirectory() as tmp:
            x = Path(tmp) / "pb.xlsx"
            pb.write_power_budget_xlsx(SUMMARY, x)
            base, _ = inp.build_base(inp.load_template(), run_mode="test", battery_wh=84.0, temperature_C=20.0)
            cfg, notes = inp.write_input(x, base, Path(tmp) / "input.json")
            self.assertEqual(json.loads((Path(tmp) / "input.json").read_text("utf-8"))["mission"]["duration_years"], 1.0)
        self.assertEqual(cfg["pack"]["cells_in_parallel"], 4)
        self.assertEqual(cfg["output"]["save_time_series"], False)
        self.assertEqual(cfg["schedule"]["sequence"][0], {"mode": "PAYLOAD", "orbits": 2})
        self.assertNotIn("SDR", cfg["timing"])                          # not a subsystem of the fixture: dropped
        self.assertTrue(any("SDR" in n for n in notes))


@unittest.skipUnless(HAS_VENDOR, "END-008 delivery not present")
class TestReader(unittest.TestCase):
    def setUp(self):
        self.s, self.o, self.ts = out.read_output(FIX / "battery_monitoring_output_small.json")
        self.summ = out.summarize(self.s, self.o)

    def test_summary_and_kpis(self):
        self.assertEqual(len(self.o), 6)
        self.assertFalse(self.summ["eol_reached"])
        ids = {k["id"]: k for k in out.kpis(self.summ)}
        self.assertEqual(set(ids), {"END-008." + k for k in (
            "soh_end_of_mission", "soh_p05", "eol_reached", "soc_min_mission", "orbits_below_soc_min",
            "law_domain_outside_fraction")})
        self.assertAlmostEqual(ids["END-008.soh_end_of_mission"]["value"], 100 * self.s["SOH_end_of_mission"]["median"], places=2)
        self.assertEqual(ids["END-008.soh_end_of_mission"]["unit"], "%")

    def test_law_domain_fraction(self):
        self.assertEqual(out.law_domain_outside_fraction({"degradation_law_domain": {
            "temperature_C": {"outside": False}, "amp": {"fraction_of_orbits_outside": 0.4},
            "cycles": {"beyond": False}}}), 0.4)
        self.assertEqual(out.law_domain_outside_fraction({"degradation_law_domain": {"temperature_C": {"outside": True}}}), 1.0)
        self.assertEqual(out.law_domain_outside_fraction({}), 0.0)

    def test_frontend_data_and_png(self):
        data = out.frontend_data(self.s, self.o, self.summ)
        self.assertEqual(data["points"], 6)
        self.assertEqual(len(data["curves"]["SOH"]), 6)
        self.assertEqual(out._decimate(5000)[0], 0)
        self.assertEqual(out._decimate(5000)[-1], 4999)
        self.assertLessEqual(len(out._decimate(5000)), out.FRONTEND_MAX_POINTS)
        with tempfile.TemporaryDirectory() as tmp:
            js = out.write_frontend_js(data, Path(tmp) / "d.js")
            self.assertTrue(js.read_text("utf-8").startswith("window.batteryHealthData = {"))
            png = out.write_png(self.s, self.o, self.ts, Path(tmp) / "a.png")
            self.assertEqual(png.read_bytes()[:4], b"\x89PNG")


@unittest.skipUnless(HAS_EXE and HAS_VENDOR and openpyxl is not None, "END-008 executable not present")
class TestExe(unittest.TestCase):
    def test_round_trip_exit_zero(self):
        with tempfile.TemporaryDirectory() as tmp:
            x = Path(tmp) / "pb.xlsx"
            pb.write_power_budget_xlsx(SUMMARY, x)
            base, _ = inp.build_base(inp.load_template(), run_mode="test", battery_wh=84.0, temperature_C=20.0)
            inp.write_input(x, base, Path(tmp) / "input.json")
            run = out.run_exe(Path(tmp) / "input.json")
            self.assertEqual(run["returncode"], 0, run["message"])
            self.assertTrue(Path(run["output_path"]).is_file())

    def test_input_errors_exit_one(self):
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "bad.json").write_text("{}", "utf-8")
            run = out.run_exe(Path(tmp) / "bad.json")
            self.assertEqual(run["returncode"], 1)
            self.assertFalse(run["ok"])

    def test_missing_exe(self):
        run = out.run_exe("x.json", vendor_dir=FIX)
        self.assertIsNone(run["returncode"])
        self.assertFalse(run["ok"])


if __name__ == "__main__":
    unittest.main()
