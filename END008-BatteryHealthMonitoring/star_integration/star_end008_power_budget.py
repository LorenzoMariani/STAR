"""STAR -> END-008: writes a power_budget.xlsx in the exact layout of the END-008 power budget from the END-004 summary.

The official converter (power_budget_to_input.py) reads the sheet by its labels and cannot read the END-004 workbook
(power_budget_tables.xlsx has another layout, a table "Number of orbits" shifted by one column and fractional orbits).
This module writes the END-008 layout (same labels, columns and formulas as power_budget.xlsx of the delivery) with numeric
input cells, from ``power_budget_summary.json`` of END-004 (key ``input_config``). Nothing here imports STAR or END-008.

Choices that the format forces (each one is returned in ``info["notes"]`` for the provenance):
- the format has ONE "Power generated per orbit": the minimum of the per-mode values of END-004 (conservative);
- the "Number of orbits" table holds integers: END-004 orbits per day by mode are rounded with the largest-remainder rule,
  so that their sum is the rounded number of orbits per day.
"""
from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Any

# END-008 sheet: mode columns IDLE, TX UHF, TX S BAND, PAYLOAD (two columns each: Sun, Eclipse) starting at column D
SHEET_MODES = [("idle", "IDLE mode"), ("tx_uhf", "TX mode UHF"), ("tx_sband", "TX mode S BAND"), ("payload", "PAYLOAD mode")]
# order of the table "Number of orbits" (columns O..R), as in the delivered example
ORBIT_TABLE = [("payload", "Payload mode"), ("tx_uhf", "TX mode UHF"), ("tx_sband", "TX mode S BAND"), ("idle", "Idle mode")]
FIRST_MODE_COL = 4                # D
ORBIT_BLOCK_COL = 14              # N (labels), O (values)


def allocate_orbits(per_day: dict[str, float]) -> dict[str, int]:
    """Integer orbits per mode (largest remainder); the total is the rounded sum of the input."""
    vals = {k: max(0.0, float(v or 0.0)) for k, v in per_day.items()}
    total = int(round(sum(vals.values())))
    out = {k: int(math.floor(v)) for k, v in vals.items()}
    missing = total - sum(out.values())
    for k in sorted(vals, key=lambda k: vals[k] - out[k], reverse=True):
        if missing <= 0:
            break
        out[k] += 1
        missing -= 1
    return out


def read_summary(summary: dict[str, Any] | str | Path) -> dict[str, Any]:
    if isinstance(summary, (str, Path)):
        summary = json.loads(Path(summary).read_text("utf-8"))
    return summary


def budget_values(summary: dict[str, Any]) -> dict[str, Any]:
    """Everything the sheet needs, from the END-004 summary, and the notes on the choices made."""
    s = read_summary(summary)
    cfg = s["input_config"]
    orbit = cfg["orbit"]
    by_mode = (cfg.get("generation") or {}).get("power_generated_per_orbit_Wh_by_mode") or {}
    gens = [float(v) for v in by_mode.values() if v is not None]
    generated = min(gens) if gens else float((cfg.get("generation") or {})["power_generated_per_orbit_Wh"])
    notes = []
    if len(set(gens)) > 1:
        notes.append("power generated per orbit = minimum over the modes ({:.4g} Wh; by mode: {}): the END-008 format has one value".format(
            generated, ", ".join("{} {:.4g}".format(k, v) for k, v in by_mode.items())))
    per_day = (cfg.get("operations") or {}).get("mode_orbits_per_day") or {}
    orbits = allocate_orbits({k: per_day.get(k, 0.0) for k, _ in SHEET_MODES})
    if any(abs(float(per_day.get(k, 0.0)) - orbits[k]) > 1e-9 for k in orbits):
        notes.append("orbits per day by mode rounded to integers: {} -> {}".format(
            {k: per_day.get(k, 0.0) for k in orbits}, orbits))
    subs = []
    for sub in cfg["subsystems"]:
        duty = {}
        for key, _ in SHEET_MODES:
            d = (sub.get("duty") or {}).get(key) or {}
            duty[key] = (float(d.get("sun", 0.0) or 0.0), float(d.get("eclipse", 0.0) or 0.0))
        subs.append({"name": str(sub["name"]), "powerW": float(sub["powerW"]), "duty": duty})
    return {"period_min": float(orbit["period_min"]), "eclipse_min": float(orbit["eclipse_min"]), "generated_Wh": generated,
            "battery_Wh": float((cfg.get("battery") or {})["capacity_Wh"]), "margin": float(cfg.get("margin_fraction", 0.1)),
            "subsystems": subs, "orbits": orbits, "notes": notes}


def write_power_budget_xlsx(summary: dict[str, Any] | str | Path, path: str | Path) -> dict[str, Any]:
    """Writes the END-008 power budget workbook; returns the values written and the notes."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter as L

    v = budget_values(summary)
    n = len(v["subsystems"])
    if n == 0:
        raise ValueError("the END-004 summary has no subsystems")
    wb = Workbook()
    ws = wb.active
    ws.title = "Power budget"
    ws["A1"], ws["B1"] = "Subsystem", "Power [W]"
    mcol = {}                                            # mode key -> column of the Sun duty (Eclipse = +1)
    for i, (key, label) in enumerate(SHEET_MODES):
        mcol[key] = FIRST_MODE_COL + 2 * i
        ws.cell(1, mcol[key], label)
    O = ORBIT_BLOCK_COL + 1
    ws.cell(1, ORBIT_BLOCK_COL, "Orbit")
    ws.cell(2, ORBIT_BLOCK_COL, "Daily orbits");       ws.cell(2, O, "=INT(1440/{}3)".format(L(O)))
    ws.cell(3, ORBIT_BLOCK_COL, "Period [min]");       ws.cell(3, O, v["period_min"])
    ws.cell(4, ORBIT_BLOCK_COL, "Sunlight [min]");     ws.cell(4, O, "={0}3-{0}5".format(L(O)))
    ws.cell(5, ORBIT_BLOCK_COL, "Eclipse [min]");      ws.cell(5, O, v["eclipse_min"])
    ws.cell(6, ORBIT_BLOCK_COL, "Power generated per orbit [W]"); ws.cell(6, O, v["generated_Wh"])
    ws.cell(7, ORBIT_BLOCK_COL, "Battery capacity [Wh]");         ws.cell(7, O, v["battery_Wh"])
    ws.cell(8, ORBIT_BLOCK_COL, "Margin");             ws.cell(8, O, v["margin"])
    sun, ecl = "${}$4".format(L(O)), "${}$5".format(L(O))
    cons_rows = []
    for i, sub in enumerate(v["subsystems"]):
        r = 2 + 3 * i
        ws.cell(r, 1, sub["name"])
        ws.cell(r, 2, sub["powerW"])
        ws.cell(r, 3, "Duty cycle - Sun/Eclipse")
        ws.cell(r + 1, 3, "Duty cycle [min]")
        ws.cell(r + 2, 3, "Consumption [Wh]")
        cons_rows.append(r + 2)
        for key, _ in SHEET_MODES:
            for k, ref in ((0, sun), (1, ecl)):
                c = mcol[key] + k
                ws.cell(r, c, round(sub["duty"][key][k], 6))
                ws.cell(r + 1, c, "={}{}*{}".format(L(c), r, ref))
                ws.cell(r + 2, c, "={}{}/60*$B${}".format(L(c), r + 1, r))
    hdr = 3 * n + 3                                      # row of the Sun/Eclipse labels (36 with 11 subsystems)
    for key, _ in SHEET_MODES:
        ws.cell(hdr, mcol[key], "Sun")
        ws.cell(hdr, mcol[key] + 1, "Eclipse")
    ws.cell(hdr + 1, 3, "Power consumption [Wh]")
    ws.cell(hdr + 2, 3, "Without margin")
    ws.cell(hdr + 3, 3, "10% margin")
    ws.cell(hdr + 4, 3, "Overall")
    ws.cell(hdr + 5, 3, "Power saved [Wh]")
    ws.cell(hdr + 6, 3, "Peak power consumption [W]")
    ws.cell(hdr + 7, 3, "Orbital DoD")
    gen, cap, mar = "${}$6".format(L(O)), "${}$7".format(L(O)), "${}$8".format(L(O))
    for key, _ in SHEET_MODES:
        s_c, e_c = mcol[key], mcol[key] + 1
        for c in (s_c, e_c):
            ws.cell(hdr + 2, c, "=" + "+".join("{}{}".format(L(c), r) for r in cons_rows))
            ws.cell(hdr + 3, c, "={0}{1}*(1+{2})".format(L(c), hdr + 2, mar))
        ws.cell(hdr + 4, s_c, "={}{}+{}{}".format(L(s_c), hdr + 3, L(e_c), hdr + 3))
        ws.cell(hdr + 5, s_c, "={}-{}{}".format(gen, L(s_c), hdr + 4))
        peaks = ["+".join("$B${}*({}{}>0)".format(2 + 3 * i, L(c), 2 + 3 * i) for i in range(n)) for c in (s_c, e_c)]
        ws.cell(hdr + 6, s_c, "=MAX({},{})".format(*peaks))
        ws.cell(hdr + 7, s_c, "=IF({0}{1}<0,-{0}{1},{2}{3})/{4}".format(L(s_c), hdr + 5, L(e_c), hdr + 2, cap))
    # table "Number of orbits" (labels one row above the values), as in the delivered sheet
    lab = 2 + 3 * (n - 1)
    ws.cell(lab + 1, ORBIT_BLOCK_COL, "Number of orbits")
    ws.cell(lab + 2, ORBIT_BLOCK_COL, "Consumption/Generation [Wh]")
    for j, (key, label) in enumerate(ORBIT_TABLE):
        c = O + j
        ws.cell(lab, c, label)
        ws.cell(lab + 1, c, int(v["orbits"][key]))
        ws.cell(lab + 2, c, "={}{}*{}{}".format(L(c), lab + 1, L(mcol[key]), hdr + 5))
    ws.cell(lab + 4, ORBIT_BLOCK_COL, "Daily margin [Wh]")
    ws.cell(lab + 4, O, "=SUM({}{}:{}{})".format(L(O), lab + 2, L(O + len(ORBIT_TABLE) - 1), lab + 2))
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions[L(ORBIT_BLOCK_COL)].width = 30
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return v
